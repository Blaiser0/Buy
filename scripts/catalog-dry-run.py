"""Local-only comparison. Reads a Supabase snapshot; NEVER writes to Supabase.

Document boundaries are reviewed paragraph offsets (zero based), not guesses
based on prices: several products have no price or embed it in a paragraph.
Descriptions remain literal source paragraphs, including embedded line breaks.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path
from collections import Counter
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports/catalog-import'
files = sorted((ROOT / 'By_photos').glob('*.docx'))
docs = [Document(p) for p in files]
db = json.loads((OUT / 'db-snapshot.json').read_text(encoding='utf-8'))

# First document's unique opening products; repeated Centellian entries use
# SKIN CARE instead. No descriptions are synthesized or merged across versions.
first = [
 (3,10,'ANUA Heartleaf Pore Control Cleansing Oil'),
 (14,22,'Anua PDRN Hyaluronic Acid Capsule 100 Serum 30ml'),
 (57,65,'Anua Niacinamide 10% + TXA 4% Serum 30ml'),
 (103,106,'Anua 10+ Azelaic Acid 10 Hyaluron Redness Soothing Serum'),
 (110,113,'ANUA 7 Rice Ceramide Hydrating Barrier Serum'),
 (116,121,'BEAUTY OF JOSEON Matte Sun Stick: Mugwort + Camelia SPF50+ PA++++'),
 (125,130,'Beauty Of Joseon - Pr Solar SPF 50+ PA ++++'),
 (134,139,'BEAUTY OF JOSEON RELIEF SUN AQUA-FRESH: RICE + B5 50ML'),
 (141,155,'CELIMAX PORE + DARK SPOT BRIGHTENING CARE SUNSCREEN'),
 (238,242,'Centellian24 Expert Madeca Cream Active Renew PDRN'),
 (257,261,'COSRX Snail 96 Mucin Power Essence'),
 (305,313,'COSRX Advanced Snail 92 All In One Cream'),
 (318,327,'K-SECRET Seoul 1988 Cleansing Oil: Pine Cica 1% + Probiotics'),
 (329,332,'K-SECRET Seoul 1988 Sun: Pine Tree + Ceramide'),
 (334,342,'K-SECRET Seoul 1988 Essence: Snail Mucin 97% + Rice'),
 (344,352,'K-SECRET Seoul 1988 Cleansing Foam: Pine Cica 1% + Probiotics'),
 (354,361,'K-SECRET SEOUL 1988 SERUM: RETINAL LIPOSOME 2% + BLACK GINSENG-serum'),
 (365,373,'K-SECRET SEOUL 1988 EYE CREAM: RETINAL LIPOSOME 4% + FERMENTED BEAN'),
]
starts = [1,10,23,26,30,41,49,56,61,78,87,100,106,113,117,120,136,142,150,158,163,166,176,184,195,200,205,210,219,226,235,239,249,256,264,269,279,282,291,296,302,308,312,322,329,341,357,362,371,376,394,400,404,418,422,432,444,454,462]
second = []
for index, start in enumerate(starts):
    end = starts[index+1] if index+1 < len(starts) else len(docs[1].paragraphs)
    name = docs[1].paragraphs[start].text.strip().split('\n')[0]
    if start < 30 and not name.lower().startswith('centellian'):
        name = 'Centellian24 ' + name
    if start in [120,136]:
        name = 'MIXSOON ' + name
    if start in [219,249]:
        name = 'TOCOBO ' + name
    if start >= 279 and 'skin1004' not in name.lower() and 'skin 1004' not in name.lower():
        name = 'SKIN1004 ' + name
    second.append((start,end,name))

def normalize(name):
    text = ''.join(c for c in unicodedata.normalize('NFD',name.lower()) if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]+', ' ',text).strip()

# Reviewed equivalences. These are ONLY used for the dry-run report, not writes.
aliases = {
 (0,125): 'Beauty of Joseon Relief Sun: Rice + Probiotics',
 (0,57): 'Anua 10+ Niacinamide 10 TXA 4 Serum',
 (0,110): 'Anua 7+ Rice Ceramide Hydrating Barrier Serum',
 (0,257): 'COSRX Advanced Snail 96 Mucin Power Essence',
 (0,354): 'K-SECRET SEOUL 1988 SERUM: Retinal liposome 2% + black ginseng',
 (1,1): 'Centellian 24+ Madeca Mela Capture Ampoule Pad',
 (1,10): 'Centellian 24 360° Shot PDRN Lifting Eye Cream',
 (1,26): 'Centellian 24+ Madeca Mela Capture Ampoule Capsule Cream',
 (1,78): 'Pure Glow Essentials Set',
 (1,106): 'Mixsoon Centella Cleansing Foam',
 (1,142): 'mixsoon Bean Cream',
 (1,235): 'TOCOBO Coconut Clay Cleansing Foam',
 (1,256): 'TOCOBO Bio Watery Sun Cream',
 (1,282): 'SKIN1004 Madagascar Centella Niacinamide 10% Boosting Shot Ampoule',
 (1,291): 'SKIN1004 Madagascar Centella Retinol 0.2% Boosting Shot Ampoule',
 (1,296): 'SKIN1004 Madagascar Centella Air-Fit Suncream Plus SPF50+ PA++++',
 (1,302): 'Skin1004 Madagascar Centella Hyalu-Cica Sleeping Pack',
 (1,308): 'Skin1004 Madagascar Centella Hyalu-Cica Blue Serum',
 (1,394): 'Skin1004 Madagascar Centella Hyalu-Cica Brightening Toner',
 (1,400): 'SKIN1004 Madagascar Centella Hyalu-Cica First Ampoule',
 (1,454): 'Skin1004 Madagascar Centella Toning Toner',
}
doubts = {
 (1,41): 'Bean Essence 30ml S/75 frente a Bean Essence sin tamaño S/89. Confirmar si son presentaciones distintas.',
 (1,113): 'Bean Essence sin tamaño S/89 frente a 30ml S/75. Confirmar contenido antes de consolidar.',
 (1,87): 'MIXSOON DOUBLE CLEANSING SET no tiene precio propio legible en el texto. S/85 pertenece al aceite siguiente.',
 (1,269): 'LATTE TIRTIR: existe Mask Fit Red Cushion (27C Cool Beige). Falta confirmar tono/presentación; no unir ni insertar.',
}

def source_text(doc, start, end):
    chunks = []
    pindex = -1
    for element in doc.element.body:
        if element.tag.endswith('}p'):
            pindex += 1
            if start <= pindex < end:
                chunks.append(Paragraph(element,doc).text)
        elif element.tag.endswith('}tbl') and start <= pindex < end:
            table = Table(element,doc)
            chunks.append('\n'.join('\t'.join(c.text for c in row.cells) for row in table.rows))
    return '\n'.join(chunks)

def category(name, text):
    n = normalize(name)
    if any(w in n for w in ['kit','set','duo','trio','suitcase','twin pack']): return 'Sets & Regalos'
    if any(w in n for w in ['sun','solar','spf']) or n == 'tocobo stick centella asiatica': return 'Protectores Solares'
    if any(w in n for w in ['cleansing','cleanser','foam']): return 'Limpiadores'
    if 'toner' in n: return 'Tónicos'
    if 'cushion' in n: return 'Maquillaje'
    if 'cream' in n: return 'Hidratantes'
    if any(w in n for w in ['serum','ampoule','essence']): return 'Sueros y Ampollas'
    return 'Hidratantes'  # Existing application/database default.

records = []
for source, specs in enumerate([first,second]):
    for start,end,name in specs:
        raw = source_text(docs[source],start,end)
        # Only standalone retail prices or prices at the end of the final line.
        prices = [m for m in re.finditer(r'S\s*/\.?\s*(\d+(?:\.\d+)?)\s*(?:SOLES)?\s*$',raw,re.I|re.M)
                  if 'recomendado' not in raw[raw.rfind('\n',0,m.start())+1:m.start()].lower()]
        price = float(prices[-1].group(1)) if prices else None
        match_name = aliases.get((source,start),name)
        matches = [p for p in db if normalize(p['name']) == normalize(match_name)]
        issue = doubts.get((source,start))
        status = 'REVISAR' if issue or len(matches)>1 else 'EXISTE' if matches else 'NUEVO'
        # Preserve raw source for audit. Do not produce an import-ready payload
        # until uncertainties and document count have been reviewed by the user.
        records.append(dict(name=name,source=files[source].name,startParagraph=start,endParagraphExclusive=end,
            sourceText=raw,price=price,category=category(name,raw),status=status,
            existing=[dict(id=p['id'],name=p['name'],price=p['price']) for p in matches],issue=issue,
            image_url=None,stock_quantity=0))

assert len({normalize(r['name']) for r in records}) == len(records), 'Repeated normalized names in manifest'
counts = Counter(r['status'] for r in records)
report = ['# Dry run de catálogo — SIN escrituras en Supabase', '',
    f'Productos solicitados: 84. Fichas candidatas distintas identificadas: {len(records)}.',
    f'Productos actuales en Supabase: {len(db)}.',
    f'Existentes: {counts["EXISTE"]}. Nuevos candidatos: {counts["NUEVO"]}. Requieren revisión: {counts["REVISAR"]}.',
    'Creados: 0. Modificados: 0. Duplicados creados: 0.', '',
    'No se confirma el total de 84: falta la lista consolidada para localizar las 7 fichas restantes o aclarar el conteo.',
    'Se conservan por separado Bean Essence 30ml y Bean Essence sin tamaño hasta aclarar su presentación.',
    'Los textos de SKIN CARE se usan como fuente para sus 59 fichas, incluyendo sus 6 tablas de modo de uso. El primer Word aporta 18 fichas adicionales.',
    'La imagen del ANUA verde permite identificar Azelaic Acid 10 Hyaluron Redness Soothing Serum; ya existe.',
    'Las imágenes incrustadas en el Word identifican ANUA PDRN Hyaluronic Acid Capsule 100 Serum 30ml y Beauty of Joseon Relief Sun Rice + Probiotics. Se inspeccionaron únicamente para identificación; NO se cargarán como imágenes de productos.',
    'El Word de Capsule Cream contiene texto que habla de Ampoule Pad. Como ya existe, se omite y no se sobrescribe.',
    'Los campos sourceText del JSON son extractos literales para auditoría, no un payload listo para importar.',
    'Antes de importar se volverá a consultar Supabase y se resolverán las coincidencias pendientes. Este script no tiene operaciones de escritura remota.', '']
for status,title in [('EXISTE','Ya existentes — se omiten'),('NUEVO','Serían creados tras aprobación'),('REVISAR','Pendientes — NO insertar')]:
    report += ['## '+title,'']
    for r in records:
        if r['status'] != status: continue
        price = f'S/ {r["price"]:.2f}' if r['price'] is not None else 'precio no identificado'
        report.append(f'- **{r["name"]}** — {price} — {r["category"]}')
        for p in r['existing']: report.append(f'  - Existe: {p["name"]} · ID {p["id"]} · precio actual S/ {p["price"]:.2f} (sin cambios).')
        if r['issue']: report.append('  - '+r['issue'])
    report += ['']
(OUT/'dry-run.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'dry-run.md').write_text('\n'.join(report),encoding='utf-8')
print('\n'.join(report))
